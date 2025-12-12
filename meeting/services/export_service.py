"""
Export Service
Exports transcripts, summaries, and action items to PDF/DOCX
"""
from typing import Dict, Any, Optional
import logging
from pathlib import Path
from datetime import datetime

from meeting.database import MeetingDatabase

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting meeting data to various formats"""
    
    def __init__(self, db: MeetingDatabase):
        """
        Initialize export service
        
        Args:
            db: MeetingDatabase instance
        """
        self.db = db
    
    def export_transcript(
        self,
        meeting_id: int,
        format: str = "pdf"
    ) -> str:
        """
        Export meeting transcript to file
        
        Args:
            meeting_id: Meeting ID
            format: Export format (pdf, docx, txt)
            
        Returns:
            Path to exported file
        """
        try:
            meeting = self.db.get_meeting(meeting_id)
            if not meeting:
                raise ValueError(f"Meeting {meeting_id} not found")
            
            segments = self.db.get_transcript(meeting_id)
            if not segments:
                raise ValueError(f"No transcript found for meeting {meeting_id}")
            
            # Combine transcript
            transcript_text = "\n".join([
                f"[{seg.get('speaker_name', 'Unknown')}]: {seg.get('text', '')}"
                for seg in segments
            ])
            
            # Generate filename
            from config import MEETING_REPORTS_FOLDER
            reports_folder = Path(MEETING_REPORTS_FOLDER)
            reports_folder.mkdir(parents=True, exist_ok=True)
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"meeting_{meeting_id}_transcript_{timestamp}.{format}"
            file_path = reports_folder / filename
            
            if format == "txt":
                return self._export_txt(file_path, meeting, transcript_text)
            elif format == "docx":
                return self._export_docx(file_path, meeting, transcript_text)
            elif format == "pdf":
                return self._export_pdf(file_path, meeting, transcript_text)
            else:
                raise ValueError(f"Unsupported format: {format}")
                
        except Exception as e:
            logger.error(f"Error exporting transcript: {e}", exc_info=True)
            raise
    
    def export_summary(
        self,
        meeting_id: int,
        format: str = "pdf"
    ) -> str:
        """
        Export meeting summary to file
        
        Args:
            meeting_id: Meeting ID
            format: Export format (pdf, docx, txt)
            
        Returns:
            Path to exported file
        """
        try:
            meeting = self.db.get_meeting(meeting_id)
            if not meeting:
                raise ValueError(f"Meeting {meeting_id} not found")
            
            summary = meeting.get("summary")
            if not summary:
                raise ValueError(f"No summary found for meeting {meeting_id}")
            
            # Generate filename
            from config import MEETING_REPORTS_FOLDER
            reports_folder = Path(MEETING_REPORTS_FOLDER)
            reports_folder.mkdir(parents=True, exist_ok=True)
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"meeting_{meeting_id}_summary_{timestamp}.{format}"
            file_path = reports_folder / filename
            
            if format == "txt":
                return self._export_txt(file_path, meeting, summary)
            elif format == "docx":
                return self._export_docx(file_path, meeting, summary)
            elif format == "pdf":
                return self._export_pdf(file_path, meeting, summary)
            else:
                raise ValueError(f"Unsupported format: {format}")
                
        except Exception as e:
            logger.error(f"Error exporting summary: {e}", exc_info=True)
            raise
    
    def export_action_items(
        self,
        meeting_id: int,
        format: str = "pdf"
    ) -> str:
        """
        Export action items to file
        
        Args:
            meeting_id: Meeting ID
            format: Export format (pdf, docx, txt)
            
        Returns:
            Path to exported file
        """
        try:
            meeting = self.db.get_meeting(meeting_id)
            if not meeting:
                raise ValueError(f"Meeting {meeting_id} not found")
            
            action_items = self.db.get_action_items(meeting_id=meeting_id)
            if not action_items:
                raise ValueError(f"No action items found for meeting {meeting_id}")
            
            # Format action items
            items_text = f"اقدامات جلسه: {meeting.get('title', 'N/A')}\n"
            items_text += f"تاریخ: {meeting.get('date', 'N/A')}\n\n"
            
            for idx, item in enumerate(action_items, 1):
                items_text += f"{idx}. {item.get('description', '')}\n"
                if item.get('assignee'):
                    items_text += f"   مسئول: {item.get('assignee')}\n"
                if item.get('due_date'):
                    items_text += f"   مهلت: {item.get('due_date')}\n"
                items_text += f"   وضعیت: {item.get('status', 'pending')}\n\n"
            
            # Generate filename
            from config import MEETING_REPORTS_FOLDER
            reports_folder = Path(MEETING_REPORTS_FOLDER)
            reports_folder.mkdir(parents=True, exist_ok=True)
            
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"meeting_{meeting_id}_action_items_{timestamp}.{format}"
            file_path = reports_folder / filename
            
            if format == "txt":
                return self._export_txt(file_path, meeting, items_text)
            elif format == "docx":
                return self._export_docx(file_path, meeting, items_text)
            elif format == "pdf":
                return self._export_pdf(file_path, meeting, items_text)
            else:
                raise ValueError(f"Unsupported format: {format}")
                
        except Exception as e:
            logger.error(f"Error exporting action items: {e}", exc_info=True)
            raise
    
    def _export_txt(self, file_path: Path, meeting: Dict[str, Any], content: str) -> str:
        """Export to TXT format"""
        try:
            header = f"جلسه: {meeting.get('title', 'N/A')}\n"
            header += f"تاریخ: {meeting.get('date', 'N/A')}\n"
            header += "=" * 50 + "\n\n"
            
            with open(file_path, "w", encoding="utf-8") as f:
                f.write(header)
                f.write(content)
            
            logger.info(f"Exported TXT to: {file_path}")
            return str(file_path)
            
        except Exception as e:
            logger.error(f"Error exporting TXT: {e}")
            raise
    
    def _export_docx(self, file_path: Path, meeting: Dict[str, Any], content: str) -> str:
        """Export to DOCX format"""
        try:
            from docx import Document
            from docx.shared import Pt
            
            doc = Document()
            
            # Add title
            title = doc.add_heading(f"جلسه: {meeting.get('title', 'N/A')}", 0)
            title.alignment = 2  # Right align
            
            # Add date
            date_para = doc.add_paragraph(f"تاریخ: {meeting.get('date', 'N/A')}")
            date_para.alignment = 2
            
            # Add separator
            doc.add_paragraph("=" * 50)
            
            # Add content
            for line in content.split('\n'):
                if line.strip():
                    para = doc.add_paragraph(line)
                    para.alignment = 2  # Right align for Persian text
            
            doc.save(str(file_path))
            logger.info(f"Exported DOCX to: {file_path}")
            return str(file_path)
            
        except ImportError:
            logger.warning("python-docx not installed, falling back to TXT")
            return self._export_txt(file_path.with_suffix('.txt'), meeting, content)
        except Exception as e:
            logger.error(f"Error exporting DOCX: {e}")
            raise
    
    def _export_pdf(self, file_path: Path, meeting: Dict[str, Any], content: str) -> str:
        """Export to PDF format"""
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.units import inch
            
            # Try to register Persian font
            try:
                from config import PERSIAN_FONT_PATH
                if Path(PERSIAN_FONT_PATH).exists():
                    pdfmetrics.registerFont(TTFont('Persian', PERSIAN_FONT_PATH))
                    font_name = 'Persian'
                else:
                    font_name = 'Helvetica'
            except Exception:
                font_name = 'Helvetica'
            
            doc = SimpleDocTemplate(
                str(file_path),
                pagesize=A4,
                rightMargin=72,
                leftMargin=72,
                topMargin=72,
                bottomMargin=18
            )
            
            story = []
            styles = getSampleStyleSheet()
            
            # Custom style for Persian text
            persian_style = ParagraphStyle(
                'PersianStyle',
                parent=styles['Normal'],
                fontName=font_name,
                fontSize=12,
                alignment=2,  # Right align
                rightIndent=0,
                leftIndent=0
            )
            
            # Add title
            title = Paragraph(f"<b>جلسه: {meeting.get('title', 'N/A')}</b>", persian_style)
            story.append(title)
            story.append(Spacer(1, 0.2*inch))
            
            # Add date
            date_para = Paragraph(f"تاریخ: {meeting.get('date', 'N/A')}", persian_style)
            story.append(date_para)
            story.append(Spacer(1, 0.2*inch))
            
            # Add separator
            story.append(Paragraph("=" * 50, persian_style))
            story.append(Spacer(1, 0.2*inch))
            
            # Add content (split into paragraphs)
            for line in content.split('\n'):
                if line.strip():
                    para = Paragraph(line.replace('<', '&lt;').replace('>', '&gt;'), persian_style)
                    story.append(para)
                    story.append(Spacer(1, 0.1*inch))
            
            doc.build(story)
            logger.info(f"Exported PDF to: {file_path}")
            return str(file_path)
            
        except ImportError:
            logger.warning("reportlab not installed, falling back to TXT")
            return self._export_txt(file_path.with_suffix('.txt'), meeting, content)
        except Exception as e:
            logger.error(f"Error exporting PDF: {e}")
            raise

